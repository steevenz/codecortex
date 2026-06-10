"""
Class OfficeWorker – Single Responsibility: Extract content from office documents.

Office document extractor integrated into graph build pipeline.

:project: CodeCortex
:package: Modules.Codegraph.Graph_builders.Office
:author: Steeven Andrian
:copyright: (c) 2026 Aegis Codework
:standard: Aegis-CodeGraph-v1.0
"""

import docx
import openpyxl
from pypdf import PdfReader
from pathlib import Path
from typing import Optional, Dict, Any

from src.core.logging import get_logger

logger = get_logger("CodeCortex.Domain.CodeGraph.OfficeWorker")

class OfficeWorker:
    """
    Production-grade extractor for Office and PDF documents.
    """
    
    def process_file(self, file_path: Path) -> Optional[str]:
        """
        Extract text from supported office formats.
        """
        ext = file_path.suffix.lower()
        
        try:
            if ext == '.docx':
                return self._extract_docx(file_path)
            elif ext == '.xlsx':
                return self._extract_xlsx(file_path)
            elif ext == '.pdf':
                return self._extract_pdf(file_path)
            elif ext in ['.doc', '.rtf']:
                # Fallback or external converter needed for legacy .doc
                logger.warning(f"Legacy format {ext} not fully supported natively yet: {file_path}")
                return None
            else:
                logger.debug(f"Unsupported office format: {ext}")
                return None
        except Exception as e:
            logger.error(f"Failed to process office file {file_path}: {e}")
            return None

    def _extract_docx(self, path: Path) -> str:
        doc = docx.Document(path)
        return "\n".join([p.text for p in doc.paragraphs])

    def _extract_xlsx(self, path: Path) -> str:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        content = []
        for sheet in wb.worksheets:
            content.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                # Filter out empty rows
                row_data = [str(cell) for cell in row if cell is not None]
                if row_data:
                    content.append("\t".join(row_data))
        return "\n".join(content)

    def _extract_pdf(self, path: Path) -> str:
        reader = PdfReader(path)
        content = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                content.append(f"--- Page {i+1} ---")
                content.append(text)
        return "\n".join(content)
